from openpyxl import Workbook
from openpyxl.chart import *
from copy import deepcopy

wb = Workbook()
ws = wb.active

rows = [('number', 'Batch1', 'Batch2'),
        (1, 10, 30),
        (2, 40, 60),
        (3, 70, 80),
        (4, 80, 90),
        (5, 95, 100),
        (6, 104, 110),
        (7, 109, 130),
        (8, 110, 140),
        (9, 117, 190)
        ]
for row in rows:
    ws.append(row)

# 图表1
chart1 = BarChart()
# 图表类型是垂直的
chart1.type = 'col'
# 设置图表样式 1-48
chart1.style = 48

# 设置图表标题
chart1.title = "Bar Chart"

# 设置X轴 Y轴的描述
chart1.x_axis.title = '月份'
chart1.y_axis.title = '销售额'
# 引用单元格的范围，用作图表展示
data = Reference(ws, min_col=2, min_row=1, max_row=8, max_col=3)
# 引用工作表的单元范围
labels = Reference(ws, min_col=2, min_row=1, max_row=8)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(labels)
ws.add_chart(chart1, 'A10')

# 第二个图表
chart2 = deepcopy(chart1)
chart2.style = 20
# 图表类型横向
chart2.type = 'bar'
chart2.title = "图表2"
ws.add_chart(chart2, 'J10')

# 图表3
chart3 = deepcopy(chart1)
chart3.type = 'col'
chart3.title = "图表3"
chart3.style = 5
chart3.grouping = 'stacked'
chart3.overlap = 100
ws.add_chart(chart3, 'A27')

# 图表4
chart4 = deepcopy(chart1)
chart4.title = '图表4'
chart4.type = 'bar'
chart4.style = 33
chart4.grouping = 'percentStacked'
chart4.overlap = 100
ws.add_chart(chart4, 'J27')

wb.save('图表.xlsx')
